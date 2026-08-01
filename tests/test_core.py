# /// script
# requires-python = ">=3.12"
# dependencies = [
#   "openpyxl>=3.1.5,<4",
#   "pandas>=2.2,<4",
#   "plotly>=6.0,<7",
#   "yfinance>=1.4,<2",
# ]
# ///
"""Focused regression tests; run with: uv run tests/test_core.py"""

from __future__ import annotations

import sys
import sqlite3
import tempfile
import unittest
from datetime import UTC, date, datetime
from pathlib import Path
from unittest.mock import patch

import pandas as pd
from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from asx_breadth.db import Database  # noqa: E402
from asx_breadth.holdings import parse_holdings  # noqa: E402
from asx_breadth.indicators import (  # noqa: E402
    AdvanceDecline,
    BenchmarkTrend,
    CurrencyIndexTrend,
    GeometricIndex,
    IndicatorContext,
    IndicatorResult,
    RatioAdjustedMcClellan,
    VolatilityTrend,
    run_indicators,
)
from asx_breadth.indicators.new_highs_lows import NewHighLow  # noqa: E402
from asx_breadth.models import (  # noqa: E402
    Holding,
    HoldingsFile,
    Snapshot,
    SnapshotInstrument,
    SnapshotSeries,
    SyncOptions,
    SyncTarget,
)
from asx_breadth.panels.charts import (  # noqa: E402
    AdvanceDeclinePanel,
    AverageCorrelationPanel,
    CurrencyIndexTrendPanel,
    PercentAboveSmaPanel,
    RealizedDispersionPanel,
    VolatilityTrendPanel,
    _ema_regime_paths,
    _signed_rasi_paths,
)
from asx_breadth.panels.summary import relative_state  # noqa: E402
from asx_breadth.providers.yahoo import YahooProvider, _normalise  # noqa: E402
from asx_breadth.sync import synchronise  # noqa: E402


class DashboardTests(unittest.TestCase):
    def test_trend_state_names_each_ema(self) -> None:
        self.assertEqual(
            relative_state(108.92, 108.97, 108.68),
            ("Below EMA19 · Above EMA39", "neutral"),
        )
        self.assertEqual(
            relative_state(110.0, 108.97, 108.68),
            ("Above EMA19 · Above EMA39", "positive"),
        )

    def test_percent_above_sma_panel_is_bounded_data_not_status(self) -> None:
        dates = pd.bdate_range("2026-07-01", periods=2)
        result = IndicatorResult(
            key="percent_above_sma",
            title="Stocks Above Moving Averages",
            frame=pd.DataFrame(
                {
                    "above_sma_20": [7, 8],
                    "eligible_sma_20": [10, 10],
                    "percent_above_sma_20": [70.0, 80.0],
                    "quality_ok": [True, True],
                },
                index=dates,
            ),
        )

        panel = PercentAboveSmaPanel(20)
        figure = panel.figure(result)
        summary = panel.summary(result)

        self.assertEqual(panel.title, "% of Stocks Above 20-Day SMA")
        self.assertEqual(summary.value, "80.0%")
        self.assertEqual(summary.detail, "8 of 10 eligible above")
        self.assertEqual([trace.name for trace in figure.data], ["% above 20SMA"])
        self.assertTrue(figure.layout.yaxis.fixedrange)
        self.assertEqual(figure.layout.yaxis.ticksuffix, "%")

    def test_realized_dispersion_panel_shows_only_dispersion_series(self) -> None:
        dates = pd.bdate_range("2026-07-01", periods=2)
        result = IndicatorResult(
            key="realized_dispersion",
            title="Realized Dispersion",
            frame=pd.DataFrame(
                {
                    "dispersion_63": [16.0, 18.5],
                    "average_stock_vol_63": [24.0, 25.5],
                    "index_vol_63": [14.0, 15.2],
                    "eligible_issues_63": [297, 298],
                    "quality_ok": [True, True],
                },
                index=dates,
            ),
        )

        panel = RealizedDispersionPanel(63)
        figure = panel.figure(result)
        summary = panel.summary(result)

        self.assertEqual(panel.title, "63-Day Realized Dispersion")
        self.assertEqual(summary.value, "18.5%")
        self.assertEqual(summary.detail, "Avg stock 25.5% · VAS 15.2% · 298 names")
        self.assertEqual([trace.name for trace in figure.data], ["Dispersion"])
        self.assertEqual(figure.layout.yaxis.ticksuffix, "%")

    def test_average_correlation_panel_shows_one_pearson_series(self) -> None:
        dates = pd.bdate_range("2026-07-01", periods=2)
        result = IndicatorResult(
            key="average_correlation",
            title="Average Stock Correlation",
            frame=pd.DataFrame(
                {
                    "average_correlation_63": [31.0, 34.5],
                    "eligible_issues_63": [297, 298],
                    "eligible_pairs_63": [43_956, 44_253],
                    "quality_ok": [True, True],
                },
                index=dates,
            ),
        )

        panel = AverageCorrelationPanel(63)
        figure = panel.figure(result)
        summary = panel.summary(result)

        self.assertEqual(panel.title, "63-Day Average Stock Correlation")
        self.assertEqual(summary.value, "34.5%")
        self.assertEqual(summary.detail, "298 stocks · 44,253 pairs")
        self.assertEqual([trace.name for trace in figure.data], ["Average correlation"])
        self.assertEqual(figure.layout.yaxis.ticksuffix, "%")


class IndicatorRunnerTests(unittest.TestCase):
    def setUp(self) -> None:
        self.context = IndicatorContext(
            snapshot=Snapshot(
                snapshot_id=1,
                universe_code="TEST",
                universe_name="Test universe",
                as_of_date=date(2026, 7, 18),
                source_path="test.xlsx",
                instruments=(),
            ),
            factors=pd.DataFrame(),
        )

    def test_resolves_dependencies_independent_of_registration_order(self) -> None:
        calls: list[str] = []
        indicators = (
            _RecordingIndicator("leaf", ("middle",), calls),
            _RecordingIndicator("root", (), calls),
            _RecordingIndicator("middle", ("root",), calls),
        )

        results = run_indicators(self.context, indicators)

        self.assertEqual(calls, ["root", "middle", "leaf"])
        self.assertEqual(list(results), calls)

    def test_rejects_duplicate_indicator_keys_before_calculation(self) -> None:
        calls: list[str] = []
        indicators = (
            _RecordingIndicator("duplicate", (), calls),
            _RecordingIndicator("duplicate", (), calls),
        )

        with self.assertRaisesRegex(ValueError, "Duplicate indicator key: 'duplicate'"):
            run_indicators(self.context, indicators)
        self.assertEqual(calls, [])

    def test_rejects_missing_dependencies_before_calculation(self) -> None:
        calls: list[str] = []

        with self.assertRaisesRegex(
            ValueError,
            r"Missing indicator dependencies: 'dependent' requires \['absent'\]",
        ):
            run_indicators(
                self.context,
                (_RecordingIndicator("dependent", ("absent",), calls),),
            )
        self.assertEqual(calls, [])

    def test_rejects_dependency_cycles_before_calculation(self) -> None:
        calls: list[str] = []
        indicators = (
            _RecordingIndicator("first", ("second",), calls),
            _RecordingIndicator("second", ("first",), calls),
        )

        with self.assertRaisesRegex(ValueError, "Cyclic indicator dependencies"):
            run_indicators(self.context, indicators)
        self.assertEqual(calls, [])


class HoldingsTests(unittest.TestCase):
    def test_finds_headers_and_maps_asx_symbols(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "holdings.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Example Fund"])
            sheet.append(["As at 30 Jun 2026"])
            sheet.append([])
            sheet.append(
                [
                    "Holding Name",
                    "Ticker",
                    "Sector",
                    "Country code",
                    "% of net assets",
                    "Market value (AUD)",
                    "# of units",
                ]
            )
            sheet.append(
                ["Commonwealth Bank", "CBA", "Banks", "AU", "10%", "$1,234", "50"]
            )
            workbook.save(path)

            parsed = parse_holdings(path)

            self.assertEqual(parsed.as_of_date, date(2026, 6, 30))
            self.assertEqual(parsed.holdings[0].provider_symbol, "CBA.AX")
            self.assertEqual(parsed.holdings[0].weight, 0.10)
            self.assertEqual(parsed.holdings[0].market_value, 1234.0)

    def test_rejects_workbook_without_an_effective_date(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "holdings.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Example Fund"])
            sheet.append(
                [
                    "Holding Name",
                    "Ticker",
                    "Sector",
                    "Country code",
                    "% of net assets",
                    "Market value (AUD)",
                    "# of units",
                ]
            )
            sheet.append(["Example", "AAA", "Banks", "AU", "1%", "$10", "1"])
            workbook.save(path)

            with self.assertRaisesRegex(ValueError, "As at"):
                parse_holdings(path)

    def test_rejects_changed_workbook_schema_with_specific_error(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "holdings.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Example Fund"])
            sheet.append(["As at 30 Jun 2026"])
            sheet.append(
                [
                    "Holding Name",
                    "Ticker",
                    "Sector",
                    "Country code",
                    "Fund percentage renamed upstream",
                    "Market value (AUD)",
                    "# of units",
                ]
            )
            sheet.append(["Example", "AAA", "Banks", "AU", "1%", "$10", "1"])
            workbook.save(path)

            with self.assertRaisesRegex(
                ValueError,
                r"schema changed or is invalid.*missing % of net assets",
            ):
                parse_holdings(path)

    def test_schema_validation_allows_missing_optional_row_details(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "holdings.xlsx"
            workbook = Workbook()
            cover = workbook.active
            cover.title = "cover"
            cover.append(["Download cover sheet"])
            sheet = workbook.create_sheet("holdings")
            sheet.append(["Example Fund"])
            sheet.append(["As at 30 Jun 2026"])
            sheet.append(
                [
                    "Ticker",
                    "Holding Name",
                    "# of units",
                    "Market value",
                    "Portfolio weight",
                    "Country",
                    "GICS sector",
                    "Unrelated extra column",
                ]
            )
            sheet.append(["AAA", "Example", None, None, None, None, None, "ignored"])
            workbook.save(path)

            parsed = parse_holdings(path)

            self.assertEqual(parsed.holdings[0].provider_symbol, "AAA.AX")
            self.assertIsNone(parsed.holdings[0].sector)
            self.assertIsNone(parsed.holdings[0].weight)
            self.assertIsNone(parsed.holdings[0].market_value)
            self.assertIsNone(parsed.holdings[0].units)

    def test_rejects_duplicate_holdings_rows(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "holdings.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Example Fund"])
            sheet.append(["As at 30 Jun 2026"])
            sheet.append(
                [
                    "Ticker",
                    "Holding Name",
                    "Sector",
                    "Country",
                    "Portfolio weight",
                    "Market value",
                    "Units",
                ]
            )
            row = ["AAA", "Example", "Banks", "AU", "1%", "$10", "1"]
            sheet.append(row)
            sheet.append(row)
            workbook.save(path)

            with self.assertRaisesRegex(ValueError, "duplicate ticker 'AAA'"):
                parse_holdings(path)

    def test_rejects_malformed_rows_inside_holdings_table(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "holdings.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["Example Fund"])
            sheet.append(["As at 30 Jun 2026"])
            sheet.append(
                [
                    "Ticker",
                    "Holding Name",
                    "Sector",
                    "Country",
                    "Portfolio weight",
                    "Market value",
                    "Units",
                ]
            )
            sheet.append(["AAA", "Example", "Banks", "AU", "not-a-number", "$10", "1"])
            workbook.save(path)

            with self.assertRaisesRegex(ValueError, "invalid portfolio weight"):
                parse_holdings(path)


class FactorCacheTests(unittest.TestCase):
    def test_schema_is_versioned_and_rejects_a_newer_cache(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "cache.sqlite3"
            database = Database(path)
            self.assertEqual(
                database.connection.execute("PRAGMA user_version").fetchone()[0],
                3,
            )
            database.close()
            connection = sqlite3.connect(path)
            connection.execute("PRAGMA user_version = 99")
            connection.commit()
            connection.close()
            with self.assertRaisesRegex(RuntimeError, "newer"):
                Database(path)

    def test_adjustment_revision_does_not_rewrite_cached_daily_factor(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            database = Database(Path(directory) / "cache.sqlite3")
            try:
                database.connection.execute(
                    """
                    INSERT INTO instruments
                        (exchange, local_symbol, provider, provider_symbol, name)
                    VALUES ('ASX', 'CBA', 'yahoo', 'CBA.AX', 'Commonwealth Bank')
                    """
                )
                instrument_id = database.connection.execute(
                    "SELECT id FROM instruments"
                ).fetchone()["id"]
                database.connection.commit()

                first_run = database.start_fetch_run(date(2026, 7, 1), date(2026, 7, 3))
                first = _frame(
                    ["2026-07-01", "2026-07-02"],
                    close=[87.0, 88.0],
                    adjusted=[75.0, 76.0],
                )
                database.append_series(
                    run_id=first_run, instrument_id=instrument_id, frame=first
                )
                database.finish_fetch_run(first_run, "complete")

                # A later dividend revision changes the displayed adjusted levels,
                # but the new day is computed from a fresh, internally consistent pair.
                second_run = database.start_fetch_run(
                    date(2026, 7, 2), date(2026, 7, 4)
                )
                revised = _frame(
                    ["2026-07-02", "2026-07-03"],
                    close=[88.0, 89.0],
                    adjusted=[74.0, 75.0],
                )
                database.append_series(
                    run_id=second_run, instrument_id=instrument_id, frame=revised
                )
                database.finish_fetch_run(second_run, "complete")

                rows = database.connection.execute(
                    """
                    SELECT trade_date, total_return_factor FROM daily_factors
                    WHERE instrument_id = ? ORDER BY trade_date
                    """,
                    (instrument_id,),
                ).fetchall()
                observations = database.connection.execute(
                    "SELECT COUNT(*) AS count FROM provider_observations"
                ).fetchone()["count"]

                self.assertEqual(len(rows), 3)
                self.assertAlmostEqual(rows[1]["total_return_factor"], 76 / 75)
                self.assertAlmostEqual(rows[2]["total_return_factor"], 75 / 74)
                self.assertEqual(observations, 4)
            finally:
                database.close()

    @patch("asx_breadth.providers.yahoo.yf.download")
    def test_yahoo_retries_an_empty_index_without_repair(
        self, download: object
    ) -> None:
        valid = _frame(
            ["2026-07-16", "2026-07-17"],
            close=[71.8, 72.3],
            adjusted=[71.8, 72.3],
        )
        download.side_effect = [pd.DataFrame(), valid]

        result = YahooProvider().fetch(
            ["^XDA"],
            start=date(2026, 7, 1),
            end=date(2026, 7, 18),
            threads=1,
            timeout=30,
        )

        self.assertEqual(len(result["^XDA"]), 2)
        self.assertTrue(download.call_args_list[0].kwargs["repair"])
        self.assertFalse(download.call_args_list[1].kwargs["repair"])

    def test_overlap_correction_revises_canonical_factors_with_audit(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            database = Database(Path(directory) / "cache.sqlite3")
            try:
                database.connection.execute(
                    """
                    INSERT INTO instruments
                        (exchange, local_symbol, provider, provider_symbol, name)
                    VALUES ('ASX', 'AAA', 'yahoo', 'AAA.AX', 'AAA')
                    """
                )
                instrument_id = database.connection.execute(
                    "SELECT id FROM instruments"
                ).fetchone()["id"]
                database.connection.commit()
                first_run = database.start_fetch_run(date(2026, 7, 1), date(2026, 7, 4))
                database.append_series(
                    run_id=first_run,
                    instrument_id=instrument_id,
                    frame=_frame(
                        ["2026-07-01", "2026-07-02", "2026-07-03"],
                        close=[100.0, 101.0, 102.0],
                        adjusted=[100.0, 101.0, 102.0],
                    ),
                )
                correction_run = database.start_fetch_run(
                    date(2026, 7, 1), date(2026, 7, 4)
                )
                database.append_series(
                    run_id=correction_run,
                    instrument_id=instrument_id,
                    frame=_frame(
                        ["2026-07-01", "2026-07-02", "2026-07-03"],
                        close=[100.0, 101.0, 102.0],
                        adjusted=[100.0, 100.0, 102.0],
                    ),
                )
                rows = database.connection.execute(
                    """
                    SELECT trade_date, total_return_factor FROM daily_factors
                    WHERE instrument_id = ? ORDER BY trade_date
                    """,
                    (instrument_id,),
                ).fetchall()
                revisions = database.connection.execute(
                    "SELECT COUNT(*) AS count FROM factor_revisions"
                ).fetchone()["count"]
                self.assertAlmostEqual(rows[1]["total_return_factor"], 1.0)
                self.assertAlmostEqual(rows[2]["total_return_factor"], 1.02)
                self.assertEqual(revisions, 2)
            finally:
                database.close()

    def test_partial_overlap_cannot_rewire_around_a_cached_date(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            database = Database(Path(directory) / "cache.sqlite3")
            try:
                database.connection.execute(
                    """
                    INSERT INTO instruments
                        (exchange, local_symbol, provider, provider_symbol, name)
                    VALUES ('ASX', 'AAA', 'yahoo', 'AAA.AX', 'AAA')
                    """
                )
                instrument_id = database.connection.execute(
                    "SELECT id FROM instruments"
                ).fetchone()["id"]
                database.connection.commit()
                first_run = database.start_fetch_run(date(2026, 7, 1), date(2026, 7, 4))
                database.append_series(
                    run_id=first_run,
                    instrument_id=instrument_id,
                    frame=_frame(
                        ["2026-07-01", "2026-07-02", "2026-07-03"],
                        close=[100.0, 101.0, 102.0],
                        adjusted=[100.0, 101.0, 102.0],
                    ),
                )
                partial_run = database.start_fetch_run(
                    date(2026, 7, 1), date(2026, 7, 4)
                )
                database.append_series(
                    run_id=partial_run,
                    instrument_id=instrument_id,
                    frame=_frame(
                        ["2026-07-01", "2026-07-03"],
                        close=[100.0, 103.0],
                        adjusted=[100.0, 103.0],
                    ),
                )
                row = database.connection.execute(
                    """
                    SELECT previous_trade_date, total_return_factor
                    FROM daily_factors
                    WHERE instrument_id = ? AND trade_date = '2026-07-03'
                    """,
                    (instrument_id,),
                ).fetchone()
                revisions = database.connection.execute(
                    "SELECT COUNT(*) FROM factor_revisions"
                ).fetchone()[0]
                self.assertEqual(row["previous_trade_date"], "2026-07-02")
                self.assertAlmostEqual(row["total_return_factor"], 102 / 101)
                self.assertEqual(revisions, 0)
            finally:
                database.close()


class MembershipRepositoryTests(unittest.TestCase):
    def test_market_series_is_registered_idempotently_and_synced(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            database = Database(Path(directory) / "cache.sqlite3")
            try:
                snapshot = database.import_snapshot(
                    _holdings_file(date(2026, 7, 1), "first", ("AAA",)),
                    universe_code="VAS",
                    universe_name="VAS",
                )
                snapshot = database.register_universe_series(
                    snapshot.snapshot_id,
                    role="benchmark",
                    provider_symbol="VAS.AX",
                    local_symbol="VAS",
                    name="VAS",
                )
                registered = database.register_universe_series(
                    snapshot.snapshot_id,
                    role="volatility",
                    provider_symbol="^AXVI",
                    local_symbol=".AXVI",
                    name="S&P/ASX 200 VIX",
                )
                repeated = database.register_universe_series(
                    snapshot.snapshot_id,
                    role="volatility",
                    provider_symbol="^AXVI",
                    local_symbol=".AXVI",
                    name="S&P/ASX 200 VIX",
                )

                volatility = repeated.instrument_for_role("volatility")
                target_symbols = {
                    target.instrument.provider_symbol
                    for target in database.sync_targets_for_snapshot(
                        snapshot.snapshot_id
                    )
                }
                self.assertIsNotNone(volatility)
                self.assertEqual(volatility.provider_symbol, "^AXVI")
                self.assertEqual(registered.market_series, repeated.market_series)
                self.assertEqual(target_symbols, {"AAA.AX", "VAS.AX", "^AXVI"})
                self.assertEqual(
                    [item.provider_symbol for item in repeated.all_instruments],
                    ["AAA.AX", "VAS.AX", "^AXVI"],
                )
            finally:
                database.close()

    def test_preserves_dated_compositions_and_syncs_transition_baskets(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            database = Database(Path(directory) / "cache.sqlite3")
            try:
                first = database.import_snapshot(
                    _holdings_file(date(2026, 7, 1), "first", ("AAA", "BBB")),
                    universe_code="VAS",
                    universe_name="VAS",
                )
                first = database.register_universe_series(
                    first.snapshot_id,
                    role="benchmark",
                    provider_symbol="VAS.AX",
                    local_symbol="VAS",
                    name="VAS",
                )
                second = database.import_snapshot(
                    _holdings_file(date(2026, 7, 8), "second", ("BBB", "CCC")),
                    universe_code="VAS",
                    universe_name="VAS",
                )

                membership = database.membership_for_snapshot(second.snapshot_id)
                compositions = {
                    effective: set(group["instrument_id"])
                    for effective, group in membership.groupby("effective_from")
                }
                first_ids = {
                    instrument.instrument_id for instrument in first.instruments
                }
                second_ids = {
                    instrument.instrument_id for instrument in second.instruments
                }
                transition_symbols = {
                    target.instrument.provider_symbol
                    for target in database.sync_targets_for_snapshot(second.snapshot_id)
                }

                self.assertEqual(compositions["2026-07-01"], first_ids)
                self.assertEqual(compositions["2026-07-08"], second_ids)
                self.assertEqual(
                    transition_symbols,
                    {"AAA.AX", "BBB.AX", "CCC.AX", "VAS.AX"},
                )

                third = database.import_snapshot(
                    _holdings_file(date(2026, 7, 15), "third", ("BBB", "CCC", "DDD")),
                    universe_code="VAS",
                    universe_name="VAS",
                )
                retry_symbols = {
                    target.instrument.provider_symbol
                    for target in database.sync_targets_for_snapshot(third.snapshot_id)
                }
                self.assertEqual(
                    retry_symbols,
                    {"AAA.AX", "BBB.AX", "CCC.AX", "DDD.AX", "VAS.AX"},
                )
                aaa_id = next(
                    instrument.instrument_id
                    for instrument in first.instruments
                    if instrument.local_symbol == "AAA"
                )
                history_start = database.prepare_sync_requirements(
                    aaa_id, calendar_days=1000
                )
                self.assertIsNotNone(history_start)
                database.mark_history_checked(aaa_id, history_start)
                database.update_sync_success(
                    instrument_id=aaa_id,
                    checked_through=date(2026, 7, 8),
                    latest_trade_date=date(2026, 7, 7),
                    backfill_attempted=False,
                    obligation_type="forward",
                )
                database.satisfy_ready_sync_requirements(aaa_id)
                still_pending = {
                    target.instrument.provider_symbol
                    for target in database.sync_targets_for_snapshot(third.snapshot_id)
                }
                self.assertIn("AAA.AX", still_pending)
                database.update_sync_success(
                    instrument_id=aaa_id,
                    checked_through=date(2026, 7, 9),
                    latest_trade_date=date(2026, 7, 8),
                    backfill_attempted=False,
                    obligation_type="forward",
                )
                database.satisfy_ready_sync_requirements(aaa_id)
                cleared_symbols = {
                    target.instrument.provider_symbol
                    for target in database.sync_targets_for_snapshot(third.snapshot_id)
                }
                self.assertEqual(
                    cleared_symbols,
                    {"BBB.AX", "CCC.AX", "DDD.AX", "VAS.AX"},
                )
            finally:
                database.close()


class FactorCacheContinuationTests(unittest.TestCase):
    def test_backward_fill_repairs_only_the_null_cache_seam(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            database = Database(Path(directory) / "cache.sqlite3")
            try:
                database.connection.execute(
                    """
                    INSERT INTO instruments
                        (exchange, local_symbol, provider, provider_symbol, name)
                    VALUES ('ASX', 'VAS', 'yahoo', 'VAS.AX', 'VAS')
                    """
                )
                instrument_id = database.connection.execute(
                    "SELECT id FROM instruments"
                ).fetchone()["id"]
                database.connection.commit()

                first_run = database.start_fetch_run(date(2026, 7, 3), date(2026, 7, 5))
                database.append_series(
                    run_id=first_run,
                    instrument_id=instrument_id,
                    frame=_frame(
                        ["2026-07-03", "2026-07-04"],
                        close=[100.0, 102.0],
                        adjusted=[90.0, 91.8],
                    ),
                )

                backfill_run = database.start_fetch_run(
                    date(2026, 7, 1), date(2026, 7, 4)
                )
                database.append_series(
                    run_id=backfill_run,
                    instrument_id=instrument_id,
                    frame=_frame(
                        ["2026-07-01", "2026-07-02", "2026-07-03"],
                        close=[96.0, 98.0, 100.0],
                        adjusted=[86.4, 88.2, 90.0],
                    ),
                )

                rows = database.connection.execute(
                    """
                    SELECT trade_date, previous_trade_date, total_return_factor
                    FROM daily_factors
                    WHERE instrument_id = ? ORDER BY trade_date
                    """,
                    (instrument_id,),
                ).fetchall()
                self.assertEqual(len(rows), 4)
                self.assertEqual(rows[2]["previous_trade_date"], "2026-07-02")
                self.assertAlmostEqual(rows[2]["total_return_factor"], 90 / 88.2)
                self.assertAlmostEqual(rows[3]["total_return_factor"], 91.8 / 90)
            finally:
                database.close()


class IndicatorTests(unittest.TestCase):
    def setUp(self) -> None:
        self.snapshot = Snapshot(
            snapshot_id=1,
            universe_code="VAS",
            universe_name="VAS",
            as_of_date=date(2026, 6, 30),
            source_path="holdings.xlsx",
            instruments=(
                SnapshotInstrument(1, "AAA", "AAA.AX", "AAA"),
                SnapshotInstrument(2, "BBB", "BBB.AX", "BBB"),
            ),
        )

    def test_geometric_index_uses_equal_dollar_total_returns_and_skips_gaps(
        self,
    ) -> None:
        dates = pd.date_range("2026-07-01", periods=4, freq="D")
        factors = pd.DataFrame(
            [
                (1, dates[0], pd.NaT, None, None),
                (1, dates[1], dates[0], 1.21, 1.44),
                (1, dates[2], dates[1], 1.10, 1.10),
                (1, dates[3], dates[2], 1.00, 1.00),
                (2, dates[0], pd.NaT, None, None),
                (2, dates[1], dates[0], 0.81, 0.64),
                (2, dates[3], dates[1], 2.00, 2.00),
            ],
            columns=[
                "instrument_id",
                "trade_date",
                "previous_trade_date",
                "close_to_previous_close",
                "total_return_factor",
            ],
        )

        result = GeometricIndex().calculate(
            IndicatorContext(
                self.snapshot,
                factors,
                series_factors={"benchmark": pd.DataFrame({"trade_date": dates})},
            ),
            {},
        )

        self.assertAlmostEqual(result.frame.iloc[0]["daily_geometric_factor"], 0.96)
        self.assertAlmostEqual(result.frame.iloc[0]["geometric_index"], 96.0)
        self.assertAlmostEqual(result.frame.iloc[1]["geometric_index"], 105.6)
        self.assertAlmostEqual(result.frame.iloc[2]["geometric_index"], 105.6)
        self.assertEqual(result.frame.iloc[2]["quoted_issues"], 1)
        self.assertIn("cash dividends", result.metadata["return_basis"])

    def test_geometric_index_accepts_a_halt_but_holds_on_broad_data_loss(self) -> None:
        dates = pd.date_range("2026-07-01", periods=3, freq="D")
        instruments = tuple(
            SnapshotInstrument(index, f"S{index}", f"S{index}.AX", f"Stock {index}")
            for index in range(1, 26)
        )
        snapshot = Snapshot(
            snapshot_id=2,
            universe_code="VAS",
            universe_name="VAS",
            as_of_date=date(2026, 6, 30),
            source_path="holdings.xlsx",
            instruments=instruments,
        )
        rows = []
        for instrument in instruments:
            rows.append((instrument.instrument_id, dates[0], pd.NaT, None))
            if instrument.instrument_id <= 24:
                rows.append((instrument.instrument_id, dates[1], dates[0], 1.01))
            if instrument.instrument_id <= 5:
                rows.append((instrument.instrument_id, dates[2], dates[1], 1.02))
        factors = pd.DataFrame(
            rows,
            columns=[
                "instrument_id",
                "trade_date",
                "previous_trade_date",
                "total_return_factor",
            ],
        )

        frame = (
            GeometricIndex()
            .calculate(
                IndicatorContext(
                    snapshot,
                    factors,
                    series_factors={"benchmark": pd.DataFrame({"trade_date": dates})},
                ),
                {},
            )
            .frame
        )

        self.assertTrue(frame.iloc[0]["quality_ok"])
        self.assertEqual(frame.iloc[0]["quoted_issues"], 24)
        self.assertAlmostEqual(frame.iloc[0]["geometric_index"], 101.0)
        self.assertFalse(frame.iloc[1]["quality_ok"])
        self.assertTrue(pd.isna(frame.iloc[1]["daily_geometric_factor"]))
        self.assertAlmostEqual(frame.iloc[1]["geometric_index"], 101.0)

    def test_ad_excludes_resume_after_missing_session(self) -> None:
        factors = pd.DataFrame(
            [
                (1, "2026-07-01", None, None),
                (1, "2026-07-02", "2026-07-01", 1.10),
                (1, "2026-07-04", "2026-07-02", 1.10),
                (2, "2026-07-01", None, None),
                (2, "2026-07-02", "2026-07-01", 0.90),
                (2, "2026-07-03", "2026-07-02", 1.00),
                (2, "2026-07-04", "2026-07-03", 1.05),
            ],
            columns=[
                "instrument_id",
                "trade_date",
                "previous_trade_date",
                "total_return_factor",
            ],
        )
        results = run_indicators(
            IndicatorContext(
                self.snapshot,
                factors,
                series_factors={
                    "benchmark": pd.DataFrame(
                        {"trade_date": pd.date_range("2026-07-01", periods=4)}
                    )
                },
            ),
            (AdvanceDecline(), RatioAdjustedMcClellan()),
        )
        ad = results["advance_decline"].frame

        self.assertEqual((ad.iloc[0]["advances"], ad.iloc[0]["declines"]), (1, 1))
        self.assertEqual(ad.iloc[1]["unchanged"], 1)
        self.assertEqual(ad.iloc[2]["advances"], 1)
        self.assertEqual(ad.iloc[2]["issues"], 1)
        self.assertAlmostEqual(
            results["mcclellan_rasi"].frame.iloc[-1]["ratio_adjusted_net_advances"],
            1000.0,
        )

    def test_ad_switches_composition_at_each_effective_date(self) -> None:
        dates = pd.date_range("2026-07-01", periods=5, freq="D")
        returns = {
            1: [None, 1.10, 1.10, 0.90, 1.10],
            2: [None, 0.90, 1.10, 0.90, 0.90],
            3: [None, 1.10, 0.90, 1.10, 0.90],
        }
        rows = []
        for instrument_id, factors in returns.items():
            for position, trade_date in enumerate(dates):
                rows.append(
                    {
                        "instrument_id": instrument_id,
                        "trade_date": trade_date,
                        "previous_trade_date": (
                            dates[position - 1] if position else pd.NaT
                        ),
                        "total_return_factor": factors[position],
                    }
                )
        membership = pd.DataFrame(
            [
                (10, "2026-07-01", 1),
                (10, "2026-07-01", 2),
                (11, "2026-07-04", 2),
                (11, "2026-07-04", 3),
            ],
            columns=["snapshot_id", "effective_from", "instrument_id"],
        )
        current = Snapshot(
            snapshot_id=11,
            universe_code="VAS",
            universe_name="VAS",
            as_of_date=date(2026, 7, 4),
            source_path="second.xlsx",
            instruments=(
                SnapshotInstrument(2, "BBB", "BBB.AX", "BBB"),
                SnapshotInstrument(3, "CCC", "CCC.AX", "CCC"),
            ),
        )

        result = (
            AdvanceDecline()
            .calculate(
                IndicatorContext(
                    snapshot=current,
                    factors=pd.DataFrame(rows),
                    membership=membership,
                    series_factors={"benchmark": pd.DataFrame({"trade_date": dates})},
                ),
                {},
            )
            .frame
        )

        self.assertEqual(result["advances"].tolist(), [1, 2, 0, 0])
        self.assertEqual(result["declines"].tolist(), [1, 0, 2, 2])
        self.assertEqual(result["cumulative_ad"].tolist(), [0, 2, 0, -2])
        self.assertTrue((result["coverage"] == 1.0).all())

    def test_cumulative_ad_includes_a_200_session_ema(self) -> None:
        dates = pd.bdate_range("2025-09-01", periods=205)
        rows = []
        for instrument_id, factor in ((1, 1.01), (2, 1.0)):
            for position, trade_date in enumerate(dates):
                rows.append(
                    {
                        "instrument_id": instrument_id,
                        "trade_date": trade_date,
                        "previous_trade_date": (
                            dates[position - 1] if position else pd.NaT
                        ),
                        "total_return_factor": factor if position else None,
                    }
                )
        result = AdvanceDecline().calculate(
            IndicatorContext(
                self.snapshot,
                pd.DataFrame(rows),
                series_factors={"benchmark": pd.DataFrame({"trade_date": dates})},
            ),
            {},
        )

        self.assertTrue(result.frame["cumulative_ad_ema200"].iloc[:199].isna().all())
        self.assertTrue(pd.notna(result.frame["cumulative_ad_ema200"].iloc[199]))
        self.assertLess(
            result.frame["cumulative_ad_ema200"].iloc[-1],
            result.frame["cumulative_ad"].iloc[-1],
        )

    def test_ad_uses_benchmark_calendar_not_anomalous_constituent_date(self) -> None:
        benchmark_dates = pd.to_datetime(["2026-07-03", "2026-07-06", "2026-07-07"])
        factors = pd.DataFrame(
            [
                (1, "2026-07-03", None, None),
                (1, "2026-07-05", "2026-07-03", 1.50),
                (1, "2026-07-06", "2026-07-03", 1.10),
                (1, "2026-07-07", "2026-07-06", 0.90),
            ],
            columns=[
                "instrument_id",
                "trade_date",
                "previous_trade_date",
                "total_return_factor",
            ],
        )
        result = (
            AdvanceDecline()
            .calculate(
                IndicatorContext(
                    snapshot=Snapshot(
                        snapshot_id=1,
                        universe_code="VAS",
                        universe_name="VAS",
                        as_of_date=date(2026, 7, 1),
                        source_path="holdings.xlsx",
                        instruments=(SnapshotInstrument(1, "AAA", "AAA.AX", "AAA"),),
                    ),
                    factors=factors,
                    series_factors={
                        "benchmark": pd.DataFrame({"trade_date": benchmark_dates})
                    },
                ),
                {},
            )
            .frame
        )
        self.assertEqual(result.index.tolist(), benchmark_dates[1:].tolist())
        self.assertEqual(result["advances"].tolist(), [1, 0])
        self.assertEqual(result["declines"].tolist(), [0, 1])

    def test_panel_allows_horizontal_but_fixes_vertical_zoom(self) -> None:
        frame = pd.DataFrame(
            {
                "advances": [1],
                "declines": [0],
                "unchanged": [0],
                "issues": [1],
                "advances_plus_declines": [1],
                "coverage": [0.5],
                "coverage_floor": [0.9],
                "quality_ok": [False],
                "cumulative_ad": [1],
                "cumulative_ad_ema19": [1],
                "cumulative_ad_ema39": [1],
                "cumulative_ad_ema200": [1],
            },
            index=pd.to_datetime(["2026-07-01"]),
        )
        result = type("Result", (), {"frame": frame})()
        figure = AdvanceDeclinePanel().figure(result)
        self.assertFalse(figure.layout.xaxis.fixedrange)
        self.assertTrue(figure.layout.yaxis.fixedrange)
        self.assertFalse(figure.layout.xaxis.rangeslider.visible)
        self.assertEqual(
            tuple(figure.layout.xaxis.rangebreaks[0].bounds),
            ("sat", "mon"),
        )
        self.assertEqual(figure.layout.dragmode, "zoom")

    def test_rasi_crossing_uses_compressed_weekday_time(self) -> None:
        frame = pd.DataFrame(
            {"rasi": [10.0, -10.0]},
            index=pd.to_datetime(["2026-07-10", "2026-07-13"]),
        )

        positive, negative = _signed_rasi_paths(frame)
        shared_zero = positive.index[
            (positive == 0) & (negative.reindex(positive.index) == 0)
        ][0]

        self.assertEqual(shared_zero.weekday(), 4)
        self.assertGreater(shared_zero, pd.Timestamp("2026-07-10"))
        self.assertLess(shared_zero, pd.Timestamp("2026-07-11"))

    def test_axvi_panel_joins_regime_lines_and_shades_to_the_ema(self) -> None:
        frame = pd.DataFrame(
            {
                "axvi": [10.0, 14.0, 8.0],
                "axvi_ema200": [12.0, 12.0, 12.0],
            },
            index=pd.to_datetime(["2026-07-10", "2026-07-13", "2026-07-14"]),
        )
        above, below, above_ema, below_ema = _ema_regime_paths(frame)
        shared = above.index[
            above.notna()
            & below.reindex(above.index).notna()
            & (above == below.reindex(above.index))
        ]
        result = IndicatorResult(
            key="volatility_trend",
            title="AXVI",
            frame=frame,
        )
        figure = VolatilityTrendPanel().figure(result)
        traces = {trace.name: trace for trace in figure.data}

        self.assertEqual(len(shared), 2)
        self.assertTrue((above_ema.loc[shared] == below_ema.loc[shared]).all())
        self.assertTrue(all(timestamp.weekday() < 5 for timestamp in shared))
        self.assertEqual(traces["AXVI above EMA"].line.color, "#b84b45")
        self.assertEqual(traces["AXVI below EMA"].line.color, "#111111")
        above_fills = [
            trace
            for trace in figure.data
            if trace.name.startswith("AXVI above EMA fill")
        ]
        below_fills = [
            trace
            for trace in figure.data
            if trace.name.startswith("AXVI below EMA fill")
        ]
        self.assertEqual(len(above_fills), 1)
        self.assertEqual(len(below_fills), 2)
        self.assertTrue(all(trace.fill == "toself" for trace in above_fills))
        self.assertTrue(all(trace.fill == "toself" for trace in below_fills))
        self.assertIsNone(traces["AXVI above EMA"].fill)
        self.assertIsNone(traces["AXVI below EMA"].fill)
        self.assertNotIn("Spread", traces["AXVI"].hovertemplate)
        self.assertNotIn("Above", traces["AXVI"].hovertemplate)
        self.assertIn("200-session EMA", traces["200-session EMA"].hovertemplate)

    def test_ratio_adjustment_divides_by_advances_plus_declines(self) -> None:
        ad = pd.DataFrame(
            {
                "advances": [4],
                "declines": [1],
                "unchanged": [5],
                "issues": [10],
                "advances_plus_declines": [5],
                "net_advances": [3],
                "quality_ok": [True],
            },
            index=pd.to_datetime(["2026-07-01"]),
        )
        result = RatioAdjustedMcClellan().calculate(
            IndicatorContext(self.snapshot, pd.DataFrame()),
            {
                "advance_decline": IndicatorResult(
                    key="advance_decline", title="A/D", frame=ad
                )
            },
        )
        self.assertAlmostEqual(
            result.frame.iloc[0]["ratio_adjusted_net_advances"], 600.0
        )
        self.assertAlmostEqual(result.frame.iloc[0]["ratio_ema19"], 60.0)
        self.assertAlmostEqual(result.frame.iloc[0]["ratio_ema39"], 30.0)
        self.assertAlmostEqual(result.frame.iloc[0]["mcclellan_oscillator"], 30.0)
        self.assertAlmostEqual(result.frame.iloc[0]["rasi"], 30.0)

    def test_rasi_holds_when_breadth_input_is_unavailable(self) -> None:
        ad = pd.DataFrame(
            {
                "advances": [4, 0],
                "declines": [1, 0],
                "advances_plus_declines": [5, 0],
                "net_advances": [3.0, float("nan")],
                "quality_ok": [True, False],
            },
            index=pd.to_datetime(["2026-07-01", "2026-07-02"]),
        )
        result = (
            RatioAdjustedMcClellan()
            .calculate(
                IndicatorContext(self.snapshot, pd.DataFrame()),
                {
                    "advance_decline": IndicatorResult(
                        key="advance_decline", title="A/D", frame=ad
                    )
                },
            )
            .frame
        )
        self.assertTrue(pd.isna(result.iloc[1]["mcclellan_oscillator"]))
        self.assertEqual(result.iloc[1]["rasi"], result.iloc[0]["rasi"])

    def test_benchmark_total_return_and_adjusted_high_low_scale(self) -> None:
        factors = pd.DataFrame(
            {
                "trade_date": pd.to_datetime(
                    ["2026-07-01", "2026-07-02", "2026-07-03"]
                ),
                "previous_trade_date": pd.to_datetime(
                    [None, "2026-07-01", "2026-07-02"]
                ),
                "total_return_factor": [None, 1.12, 1.02],
                "close_to_previous_close": [None, 1.10, 1.01],
                "high_to_previous_close": [None, 1.15, 1.04],
                "low_to_previous_close": [None, 1.05, 0.99],
            }
        )
        benchmark = SnapshotInstrument(3, "VAS", "VAS.AX", "VAS")
        snapshot = Snapshot(
            snapshot_id=1,
            universe_code="VAS",
            universe_name="VAS",
            as_of_date=date(2026, 6, 30),
            source_path="holdings.xlsx",
            instruments=(),
            market_series=(SnapshotSeries("benchmark", benchmark),),
        )
        result = (
            BenchmarkTrend()
            .calculate(
                IndicatorContext(
                    snapshot,
                    pd.DataFrame(),
                    series_factors={"benchmark": factors},
                    series_anchor_prices={"benchmark": 57.12},
                ),
                {},
            )
            .frame
        )
        self.assertAlmostEqual(result.iloc[-1]["total_return_index"], 57.12)
        self.assertAlmostEqual(result.iloc[1]["total_return_index"], 56.0)
        self.assertAlmostEqual(
            result.iloc[1]["adjusted_high_index"],
            0.5 * 100 * 1.15 * 1.12 / 1.10,
        )
        self.assertTrue(result["low_ema200"].isna().all())
        self.assertTrue(result["high_ema200"].isna().all())

    def test_axvi_reconstructs_actual_level_backwards_from_cached_anchor(self) -> None:
        dates = pd.bdate_range("2025-09-01", periods=205)
        factors = pd.DataFrame(
            {
                "trade_date": dates,
                "previous_trade_date": [pd.NaT, *dates[:-1]],
                "total_return_factor": [None, *([1.01] * 204)],
            }
        )
        volatility = SnapshotInstrument(
            99,
            ".AXVI",
            "^AXVI",
            "S&P/ASX 200 VIX",
        )
        snapshot = Snapshot(
            snapshot_id=1,
            universe_code="VAS",
            universe_name="VAS",
            as_of_date=date(2026, 6, 30),
            source_path="holdings.xlsx",
            instruments=(),
            market_series=(SnapshotSeries("volatility", volatility),),
        )

        result = VolatilityTrend().calculate(
            IndicatorContext(
                snapshot=snapshot,
                factors=pd.DataFrame(),
                series_factors={"volatility": factors},
                series_anchor_prices={"volatility": 24.5},
            ),
            {},
        )

        self.assertAlmostEqual(result.frame.iloc[-1]["axvi"], 24.5)
        self.assertAlmostEqual(
            result.frame.iloc[0]["axvi"],
            24.5 / (1.01**204),
        )
        self.assertTrue(pd.notna(result.frame.iloc[-1]["axvi_ema200"]))
        self.assertEqual(result.metadata["provider_symbol"], "^AXVI")

    def test_xda_uses_actual_level_with_19_39_and_200_session_emas(self) -> None:
        dates = pd.bdate_range("2025-09-01", periods=205)
        factors = pd.DataFrame(
            {
                "trade_date": dates,
                "previous_trade_date": [pd.NaT, *dates[:-1]],
                "total_return_factor": [None, *([1.001] * 204)],
            }
        )
        currency_index = SnapshotInstrument(
            100,
            "XDA",
            "^XDA",
            "Australian Dollar Currency Index",
        )
        snapshot = Snapshot(
            snapshot_id=1,
            universe_code="VAS",
            universe_name="VAS",
            as_of_date=date(2026, 6, 30),
            source_path="holdings.xlsx",
            instruments=(),
            market_series=(SnapshotSeries("currency_index", currency_index),),
        )

        result = CurrencyIndexTrend().calculate(
            IndicatorContext(
                snapshot=snapshot,
                factors=pd.DataFrame(),
                series_factors={"currency_index": factors},
                series_anchor_prices={"currency_index": 72.3},
            ),
            {},
        )
        figure = CurrencyIndexTrendPanel().figure(result)

        self.assertAlmostEqual(result.frame.iloc[-1]["currency_index"], 72.3)
        self.assertTrue(pd.notna(result.frame.iloc[-1]["currency_index_ema19"]))
        self.assertTrue(pd.notna(result.frame.iloc[-1]["currency_index_ema39"]))
        self.assertTrue(pd.notna(result.frame.iloc[-1]["currency_index_ema200"]))
        self.assertEqual(
            {trace.name for trace in figure.data},
            {"XDA", "19-session EMA", "39-session EMA", "200-session EMA"},
        )

    def test_new_high_low_requires_full_lookback_and_uses_price_extremes(self) -> None:
        rows = []
        dates = pd.date_range("2026-07-01", periods=5, freq="D")
        for instrument_id, factor in ((1, 1.10), (2, 0.90)):
            for position, trade_date in enumerate(dates):
                rows.append(
                    {
                        "instrument_id": instrument_id,
                        "trade_date": trade_date,
                        "previous_trade_date": (
                            dates[position - 1] if position else pd.NaT
                        ),
                        "close_to_previous_close": factor if position else None,
                        "total_return_factor": factor if position else None,
                        "high_to_previous_close": (
                            (1.10 if instrument_id == 1 else 1.00) if position else None
                        ),
                        "low_to_previous_close": (
                            (1.00 if instrument_id == 1 else 0.90) if position else None
                        ),
                    }
                )
        frame = (
            NewHighLow(lookback_sessions=3)
            .calculate(
                IndicatorContext(
                    self.snapshot,
                    pd.DataFrame(rows),
                    series_factors={"benchmark": pd.DataFrame({"trade_date": dates})},
                ),
                {},
            )
            .frame
        )
        self.assertEqual(frame.index.min(), dates[3])
        self.assertTrue((frame["new_highs"] == 1).all())
        self.assertTrue((frame["new_lows"] == 1).all())
        self.assertTrue((frame["nh_nl"] == 0).all())

    def test_new_high_low_uses_pre_entry_prices_but_only_counts_active_member(
        self,
    ) -> None:
        rows = []
        dates = pd.date_range("2026-07-01", periods=5, freq="D")
        for instrument_id, factor in ((1, 1.10), (2, 0.90)):
            for position, trade_date in enumerate(dates):
                rows.append(
                    {
                        "instrument_id": instrument_id,
                        "trade_date": trade_date,
                        "previous_trade_date": (
                            dates[position - 1] if position else pd.NaT
                        ),
                        "close_to_previous_close": factor if position else None,
                        "total_return_factor": factor if position else None,
                        "high_to_previous_close": (
                            (1.10 if instrument_id == 1 else 1.00) if position else None
                        ),
                        "low_to_previous_close": (
                            (1.00 if instrument_id == 1 else 0.90) if position else None
                        ),
                    }
                )
        membership = pd.DataFrame(
            [
                (10, "2026-07-01", 1),
                (11, "2026-07-04", 2),
            ],
            columns=["snapshot_id", "effective_from", "instrument_id"],
        )
        current = Snapshot(
            snapshot_id=11,
            universe_code="VAS",
            universe_name="VAS",
            as_of_date=date(2026, 7, 4),
            source_path="second.xlsx",
            instruments=(SnapshotInstrument(2, "BBB", "BBB.AX", "BBB"),),
        )

        frame = (
            NewHighLow(lookback_sessions=3)
            .calculate(
                IndicatorContext(
                    snapshot=current,
                    factors=pd.DataFrame(rows),
                    membership=membership,
                    series_factors={"benchmark": pd.DataFrame({"trade_date": dates})},
                ),
                {},
            )
            .frame
        )

        self.assertEqual(frame.index.min(), dates[3])
        self.assertEqual(frame["new_highs"].tolist(), [1.0, 0.0])
        self.assertEqual(frame["new_lows"].tolist(), [0.0, 1.0])

    def test_new_high_low_tolerates_an_isolated_history_gap(self) -> None:
        dates = pd.date_range("2026-07-01", periods=7, freq="D")
        rows = []
        for position, trade_date in enumerate(dates):
            if position == 3:
                continue
            rows.append(
                {
                    "instrument_id": 1,
                    "trade_date": trade_date,
                    "previous_trade_date": (
                        dates[position - 2]
                        if position == 4
                        else (dates[position - 1] if position else pd.NaT)
                    ),
                    "close_to_previous_close": 1.10 if position else None,
                    "total_return_factor": 1.10 if position else None,
                    "high_to_previous_close": 1.10 if position else None,
                    "low_to_previous_close": 1.00 if position else None,
                }
            )
        current = Snapshot(
            snapshot_id=1,
            universe_code="VAS",
            universe_name="VAS",
            as_of_date=date(2026, 7, 1),
            source_path="holdings.xlsx",
            instruments=(SnapshotInstrument(1, "AAA", "AAA.AX", "AAA"),),
        )
        frame = (
            NewHighLow(
                lookback_sessions=5,
                minimum_history_coverage=0.75,
            )
            .calculate(
                IndicatorContext(
                    snapshot=current,
                    factors=pd.DataFrame(rows),
                    series_factors={"benchmark": pd.DataFrame({"trade_date": dates})},
                ),
                {},
            )
            .frame
        )
        self.assertEqual(frame.index[-1], dates[-1])
        self.assertEqual(frame.iloc[-1]["new_highs"], 1)

    def test_new_high_low_withholds_a_partial_broad_universe_session(self) -> None:
        dates = pd.date_range("2026-07-01", periods=5, freq="D")
        instruments = tuple(
            SnapshotInstrument(index, f"S{index}", f"S{index}.AX", f"S{index}")
            for index in range(1, 21)
        )
        rows = []
        for instrument in instruments:
            for position, trade_date in enumerate(dates):
                if position == len(dates) - 1 and instrument.instrument_id > 15:
                    continue
                rows.append(
                    {
                        "instrument_id": instrument.instrument_id,
                        "trade_date": trade_date,
                        "previous_trade_date": (
                            dates[position - 1] if position else pd.NaT
                        ),
                        "close_to_previous_close": 1.01 if position else None,
                        "total_return_factor": 1.01 if position else None,
                        "high_to_previous_close": 1.01 if position else None,
                        "low_to_previous_close": 1.00 if position else None,
                    }
                )
        snapshot = Snapshot(
            snapshot_id=1,
            universe_code="TEST",
            universe_name="TEST",
            as_of_date=date(2026, 7, 1),
            source_path="holdings.xlsx",
            instruments=instruments,
        )
        frame = (
            NewHighLow(lookback_sessions=3)
            .calculate(
                IndicatorContext(
                    snapshot=snapshot,
                    factors=pd.DataFrame(rows),
                    series_factors={"benchmark": pd.DataFrame({"trade_date": dates})},
                ),
                {},
            )
            .frame
        )
        self.assertAlmostEqual(frame.iloc[-1]["coverage"], 0.75)
        self.assertFalse(frame.iloc[-1]["quality_ok"])
        self.assertTrue(pd.isna(frame.iloc[-1]["nh_nl"]))


class ProviderAndSyncTests(unittest.TestCase):
    def test_sync_options_reject_invalid_limits(self) -> None:
        with self.assertRaisesRegex(ValueError, "must be positive"):
            SyncOptions(retries=0)
        with self.assertRaisesRegex(ValueError, "must be non-negative"):
            SyncOptions(batch_pause=-0.1)

    def test_missing_adjusted_close_is_rejected(self) -> None:
        raw = pd.DataFrame(
            {
                "Open": [10.0],
                "High": [11.0],
                "Low": [9.0],
                "Close": [10.0],
                "Volume": [100.0],
            },
            index=pd.to_datetime(["2026-07-01"]),
        )
        self.assertTrue(
            _normalise(raw, start=date(2026, 7, 1), end=date(2026, 7, 2)).empty
        )

    def test_missing_backfill_anchor_does_not_advance_watermark(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            database = Database(Path(directory) / "cache.sqlite3")
            try:
                database.connection.execute(
                    """
                    INSERT INTO instruments
                        (exchange, local_symbol, provider, provider_symbol, name)
                    VALUES ('ASX', 'AAA', 'yahoo', 'AAA.AX', 'AAA')
                    """
                )
                instrument_id = database.connection.execute(
                    "SELECT id FROM instruments"
                ).fetchone()["id"]
                database.connection.commit()
                run_id = database.start_fetch_run(date(2026, 7, 3), date(2026, 7, 5))
                database.append_series(
                    run_id=run_id,
                    instrument_id=instrument_id,
                    frame=_frame(
                        ["2026-07-03", "2026-07-04"],
                        close=[100.0, 101.0],
                        adjusted=[100.0, 101.0],
                    ),
                )
                database.update_sync_success(
                    instrument_id=instrument_id,
                    checked_through=date(2026, 7, 11),
                    latest_trade_date=date(2026, 7, 4),
                    backfill_attempted=True,
                )
                snapshot = Snapshot(
                    snapshot_id=1,
                    universe_code="TEST",
                    universe_name="TEST",
                    as_of_date=date(2026, 7, 1),
                    source_path="holdings.xlsx",
                    instruments=(SnapshotInstrument(1, "AAA", "AAA.AX", "AAA"),),
                )
                provider = _FakeProvider(
                    _frame(
                        ["2026-05-20", "2026-05-21"],
                        close=[90.0, 91.0],
                        adjusted=[90.0, 91.0],
                    )
                )
                report = synchronise(
                    database,
                    _active_targets(snapshot),
                    SyncOptions(
                        lookback_sessions=10,
                        batch_size=10,
                        retries=1,
                        batch_pause=0,
                        base_backoff=0,
                    ),
                    now=datetime(2026, 7, 10, 8, tzinfo=UTC),
                    provider=provider,
                )
                self.assertEqual(report.failed, 1)
                self.assertEqual(database.history_watermarks([instrument_id]), {})
            finally:
                database.close()

    def test_anchor_invalid_response_is_retried_before_state_advances(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            database = Database(Path(directory) / "cache.sqlite3")
            try:
                database.connection.execute(
                    """
                    INSERT INTO instruments
                        (exchange, local_symbol, provider, provider_symbol, name)
                    VALUES ('ASX', 'AAA', 'yahoo', 'AAA.AX', 'AAA')
                    """
                )
                instrument_id = database.connection.execute(
                    "SELECT id FROM instruments"
                ).fetchone()["id"]
                database.connection.commit()
                run_id = database.start_fetch_run(date(2026, 7, 3), date(2026, 7, 5))
                database.append_series(
                    run_id=run_id,
                    instrument_id=instrument_id,
                    frame=_frame(
                        ["2026-07-03", "2026-07-04"],
                        close=[100.0, 101.0],
                        adjusted=[100.0, 101.0],
                    ),
                )
                database.update_sync_success(
                    instrument_id=instrument_id,
                    checked_through=date(2026, 7, 11),
                    latest_trade_date=date(2026, 7, 4),
                    backfill_attempted=True,
                )
                database.mark_history_checked(instrument_id, date(2026, 5, 1))
                target = SyncTarget(
                    SnapshotInstrument(1, "AAA", "AAA.AX", "AAA"),
                    active=True,
                )
                provider = _SequencedProvider(
                    [
                        _frame(
                            ["2026-07-10", "2026-07-11"],
                            close=[102.0, 103.0],
                            adjusted=[102.0, 103.0],
                        ),
                        _frame(
                            ["2026-07-03", "2026-07-04", "2026-07-17"],
                            close=[100.0, 101.0, 104.0],
                            adjusted=[100.0, 101.0, 104.0],
                        ),
                    ]
                )
                report = synchronise(
                    database,
                    (target,),
                    SyncOptions(
                        lookback_sessions=10,
                        batch_size=10,
                        retries=2,
                        batch_pause=0,
                        base_backoff=0,
                    ),
                    now=datetime(2026, 7, 17, 8, tzinfo=UTC),
                    provider=provider,
                )
                self.assertEqual(provider.call_count, 2)
                # The second response repairs the missing anchor, but its
                # otherwise sparse horizon remains provisional until a
                # separate fetch run repeats the same evidence.
                self.assertEqual(report.successful, 0)
                self.assertEqual(report.failed, 1)
                confirmation = synchronise(
                    database,
                    (target,),
                    SyncOptions(
                        lookback_sessions=10,
                        batch_size=10,
                        retries=1,
                        batch_pause=0,
                        base_backoff=0,
                    ),
                    now=datetime(2026, 7, 18, 8, tzinfo=UTC),
                    provider=_FakeProvider(
                        _frame(
                            ["2026-07-03", "2026-07-04", "2026-07-17"],
                            close=[100.0, 101.0, 104.0],
                            adjusted=[100.0, 101.0, 104.0],
                        )
                    ),
                )
                self.assertEqual(confirmation.successful, 1)
                self.assertEqual(confirmation.failed, 0)
            finally:
                database.close()

    def test_outgoing_target_fetch_stops_at_its_membership_boundary(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            database = Database(Path(directory) / "cache.sqlite3")
            try:
                database.connection.execute(
                    """
                    INSERT INTO instruments
                        (exchange, local_symbol, provider, provider_symbol, name)
                    VALUES ('ASX', 'AAA', 'yahoo', 'AAA.AX', 'AAA')
                    """
                )
                instrument_id = database.connection.execute(
                    "SELECT id FROM instruments"
                ).fetchone()["id"]
                database.connection.commit()
                run_id = database.start_fetch_run(date(2026, 7, 1), date(2026, 7, 3))
                database.append_series(
                    run_id=run_id,
                    instrument_id=instrument_id,
                    frame=_frame(
                        ["2026-07-01", "2026-07-02"],
                        close=[100.0, 101.0],
                        adjusted=[100.0, 101.0],
                    ),
                )
                database.update_sync_success(
                    instrument_id=instrument_id,
                    checked_through=date(2026, 7, 3),
                    latest_trade_date=date(2026, 7, 2),
                    backfill_attempted=True,
                )
                provider = _FakeProvider(
                    _frame(
                        ["2026-07-01", "2026-07-02", "2026-07-08"],
                        close=[100.0, 101.0, 102.0],
                        adjusted=[100.0, 101.0, 102.0],
                    )
                )
                report = synchronise(
                    database,
                    (
                        SyncTarget(
                            SnapshotInstrument(
                                instrument_id,
                                "AAA",
                                "AAA.AX",
                                "AAA",
                            ),
                            active=False,
                            required_end=date(2026, 7, 8),
                        ),
                    ),
                    SyncOptions(
                        lookback_sessions=10,
                        batch_size=10,
                        retries=1,
                        batch_pause=0,
                        base_backoff=0,
                    ),
                    now=datetime(2026, 7, 17, 8, tzinfo=UTC),
                    provider=provider,
                )
                self.assertEqual(report.successful, 1)
                self.assertEqual(provider.calls[0][1], date(2026, 7, 9))
            finally:
                database.close()

    def test_truncated_initial_history_needs_a_second_matching_probe(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            database = Database(Path(directory) / "cache.sqlite3")
            try:
                snapshot = database.import_snapshot(
                    _holdings_file(date(2026, 7, 1), "first", ("AAA",)),
                    universe_code="TEST",
                    universe_name="TEST",
                )
                provider = _FakeProvider(
                    _frame(
                        ["2026-06-20", "2026-06-21"],
                        close=[90.0, 91.0],
                        adjusted=[90.0, 91.0],
                    )
                )
                options = SyncOptions(
                    lookback_sessions=10,
                    batch_size=10,
                    retries=1,
                    batch_pause=0,
                    base_backoff=0,
                )
                now = datetime(2026, 7, 10, 8, tzinfo=UTC)

                synchronise(
                    database,
                    (SyncTarget(snapshot.instruments[0], active=True),),
                    options,
                    now=now,
                    provider=provider,
                )
                instrument_id = snapshot.instruments[0].instrument_id
                self.assertEqual(database.history_watermarks([instrument_id]), {})

                second_report = synchronise(
                    database,
                    (SyncTarget(snapshot.instruments[0], active=True),),
                    options,
                    now=datetime(2026, 7, 17, 8, tzinfo=UTC),
                    provider=provider,
                )
                watermark = database.history_watermarks([instrument_id])[instrument_id]
                self.assertLessEqual(watermark, date(2026, 5, 13))
                self.assertEqual(second_report.requested, 2)
                self.assertEqual(len(provider.calls), 3)
                self.assertTrue(
                    any(end == date(2026, 7, 18) for _, end in provider.calls)
                )
            finally:
                database.close()


def _frame(
    dates: list[str], *, close: list[float], adjusted: list[float]
) -> pd.DataFrame:
    return pd.DataFrame(
        {
            "open": close,
            "high": close,
            "low": close,
            "close": close,
            "adjusted_close": adjusted,
            "volume": [100.0] * len(dates),
            "dividend": [0.0] * len(dates),
            "split_ratio": [0.0] * len(dates),
            "repaired": [False] * len(dates),
        },
        index=pd.to_datetime(dates),
    )


def _holdings_file(
    as_of_date: date,
    sha256: str,
    symbols: tuple[str, ...],
) -> HoldingsFile:
    return HoldingsFile(
        path=Path(f"{sha256}.xlsx"),
        as_of_date=as_of_date,
        fund_name="VAS",
        holdings=tuple(
            Holding(
                local_symbol=symbol,
                provider_symbol=f"{symbol}.AX",
                name=symbol,
                sector=None,
                country_code="AU",
                weight=None,
                market_value=None,
                units=None,
            )
            for symbol in symbols
        ),
        sha256=sha256,
    )


class _RecordingIndicator:
    def __init__(
        self,
        key: str,
        dependencies: tuple[str, ...],
        calls: list[str],
    ) -> None:
        self.key = key
        self.dependencies = dependencies
        self.calls = calls

    def calculate(
        self,
        context: IndicatorContext,
        results: dict[str, IndicatorResult],
    ) -> IndicatorResult:
        del context
        assert all(dependency in results for dependency in self.dependencies)
        self.calls.append(self.key)
        return IndicatorResult(self.key, self.key, pd.DataFrame())


class _FakeProvider:
    def __init__(self, frame: pd.DataFrame):
        self.frame = frame
        self.calls: list[tuple[date, date]] = []

    def fetch(
        self,
        symbols: list[str],
        *,
        start: date,
        end: date,
        **_: object,
    ) -> dict[str, pd.DataFrame]:
        self.calls.append((start, end))
        return {symbol: self.frame.copy() for symbol in symbols}


class _SequencedProvider:
    def __init__(self, frames: list[pd.DataFrame]):
        self.frames = frames
        self.call_count = 0

    def fetch(self, symbols: list[str], **_: object) -> dict[str, pd.DataFrame]:
        position = min(self.call_count, len(self.frames) - 1)
        frame = self.frames[position]
        self.call_count += 1
        return {symbol: frame.copy() for symbol in symbols}


def _active_targets(snapshot: Snapshot) -> tuple[SyncTarget, ...]:
    return tuple(
        SyncTarget(instrument=instrument, active=True)
        for instrument in snapshot.all_instruments
    )


if __name__ == "__main__":
    unittest.main()
