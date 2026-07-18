"""Import Vanguard-style holdings workbooks without fixed row numbers."""

from __future__ import annotations

import hashlib
import math
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .models import Holding, HoldingsFile


HEADER_ALIASES = {
    "ticker": {"ticker", "symbol", "asx code", "asx ticker"},
    "name": {"holding name", "name", "security name"},
    "sector": {"sector", "gics sector"},
    "country": {"country code", "country"},
    "weight": {"% of net assets", "weight", "portfolio weight"},
    "market_value": {"market value (aud)", "market value"},
    "units": {"# of units", "units", "shares"},
}

HEADER_LABELS = {
    "ticker": "Ticker",
    "name": "Holding Name",
    "sector": "Sector",
    "country": "Country code",
    "weight": "% of net assets",
    "market_value": "Market value (AUD)",
    "units": "# of units",
}


class HoldingsSchemaError(ValueError):
    """The workbook structure cannot be mapped without guessing."""


def _normalise_header(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().lower())


def _column_matches(values: tuple[Any, ...]) -> dict[str, list[int]]:
    normalised = [_normalise_header(value) for value in values]
    result: dict[str, list[int]] = {}
    for key, aliases in HEADER_ALIASES.items():
        result[key] = [
            index for index, value in enumerate(normalised) if value in aliases
        ]
    return result


def _find_holdings_table(
    sheets: list[tuple[str, list[tuple[Any, ...]]]],
) -> tuple[str, list[tuple[Any, ...]], int, dict[str, int]]:
    candidates: list[
        tuple[int, str, list[tuple[Any, ...]], int, dict[str, list[int]]]
    ] = []
    for sheet_name, rows in sheets:
        for row_index, row in enumerate(rows[:50]):
            matches = _column_matches(row)
            score = sum(bool(indices) for indices in matches.values())
            if score >= 2:
                candidates.append((score, sheet_name, rows, row_index, matches))

    complete = [
        candidate
        for candidate in candidates
        if all(len(indices) == 1 for indices in candidate[4].values())
    ]
    if len(complete) > 1:
        locations = ", ".join(
            f"{sheet_name!r} row {row_index + 1}"
            for _, sheet_name, _, row_index, _ in complete
        )
        raise HoldingsSchemaError(
            "Holdings workbook schema is ambiguous: multiple complete holdings "
            f"tables were found at {locations}"
        )
    if len(complete) == 1:
        _, sheet_name, rows, row_index, matches = complete[0]
        return (
            sheet_name,
            rows,
            row_index,
            {key: indices[0] for key, indices in matches.items()},
        )

    expected = ", ".join(HEADER_LABELS.values())
    if not candidates:
        raise HoldingsSchemaError(
            "Holdings workbook schema is invalid or unsupported: could not find a "
            f"header row containing the expected columns ({expected})"
        )

    _, sheet_name, _, row_index, matches = max(candidates, key=lambda item: item[0])
    missing = [HEADER_LABELS[key] for key, indices in matches.items() if not indices]
    ambiguous = [
        f"{HEADER_LABELS[key]} in columns "
        + "/".join(get_column_letter(index + 1) for index in indices)
        for key, indices in matches.items()
        if len(indices) > 1
    ]
    problems: list[str] = []
    if missing:
        problems.append("missing " + ", ".join(missing))
    if ambiguous:
        problems.append("ambiguous " + ", ".join(ambiguous))
    raise HoldingsSchemaError(
        f"Holdings workbook schema changed or is invalid on sheet {sheet_name!r}, "
        f"row {row_index + 1}: {'; '.join(problems)}"
    )


def _number(value: Any, *, percentage: bool = False) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        number = float(value)
        if not math.isfinite(number):
            return None
        return number if not percentage or number <= 1 else number / 100
    text = str(value).strip()
    is_percent = text.endswith("%")
    text = re.sub(r"[$,%\s]", "", text)
    try:
        number = float(text)
    except ValueError:
        return None
    if not math.isfinite(number):
        return None
    return number / 100 if percentage and (is_percent or number > 1) else number


def _holding_number(
    value: Any,
    *,
    label: str,
    row_number: int,
    percentage: bool = False,
) -> float | None:
    parsed = _number(value, percentage=percentage)
    if parsed is None and value is not None and str(value).strip():
        raise HoldingsSchemaError(
            f"Holdings workbook has invalid {label} {value!r} on row {row_number}"
        )
    return parsed


def _as_of_date(rows: list[tuple[Any, ...]]) -> date | None:
    patterns = ("%d %b %Y", "%d %B %Y", "%Y-%m-%d", "%d/%m/%Y")
    for row in rows:
        for value in row:
            match = re.search(r"\b(?:as at|as of)\s+(.+)$", str(value or ""), re.I)
            if not match:
                continue
            text = match.group(1).strip()
            for pattern in patterns:
                try:
                    return datetime.strptime(text, pattern).date()
                except ValueError:
                    pass
    return None


def yahoo_symbol(local_symbol: str, country_code: str | None) -> str:
    """Map an ASX local code to Yahoo's exchange-qualified symbol."""
    symbol = local_symbol.strip().upper()
    if symbol.endswith(".AX"):
        return symbol
    if country_code in (None, "", "AU") and "." not in symbol:
        return f"{symbol}.AX"
    return symbol


def _holding_from_row(
    row: tuple[Any, ...],
    columns: dict[str, int],
    *,
    row_number: int,
) -> Holding:
    def value(key: str) -> Any:
        position = columns[key]
        return row[position] if position < len(row) else None

    ticker = str(value("ticker") or "").strip().upper()
    name = str(value("name") or "").strip()
    if not ticker or not name or not re.fullmatch(r"[A-Z0-9.-]+", ticker):
        raise HoldingsSchemaError(
            "Holdings workbook has an invalid holding row at "
            f"{row_number}: expected a name and ASX ticker"
        )
    country = str(value("country") or "").strip().upper() or None
    return Holding(
        local_symbol=ticker,
        provider_symbol=yahoo_symbol(ticker, country),
        name=name,
        sector=str(value("sector") or "").strip() or None,
        country_code=country,
        weight=_holding_number(
            value("weight"),
            label="portfolio weight",
            row_number=row_number,
            percentage=True,
        ),
        market_value=_holding_number(
            value("market_value"),
            label="market value",
            row_number=row_number,
        ),
        units=_holding_number(
            value("units"),
            label="unit count",
            row_number=row_number,
        ),
    )


def parse_holdings(
    path: Path,
    *,
    as_of_date_override: date | None = None,
) -> HoldingsFile:
    path = path.expanduser().resolve()
    if not path.is_file():
        raise ValueError(f"Holdings workbook does not exist: {path}")

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise HoldingsSchemaError(
            f"Could not read holdings workbook {path.name!r}: {exc}"
        ) from exc
    try:
        sheets = [
            (
                sheet.title,
                [tuple(row) for row in sheet.iter_rows(values_only=True)],
            )
            for sheet in workbook.worksheets
        ]
    finally:
        workbook.close()

    _, rows, header_index, columns = _find_holdings_table(sheets)

    fund_name = str(rows[0][0] or path.stem).strip()
    holdings: list[Holding] = []
    seen: dict[str, int] = {}
    for row_number, row in enumerate(
        rows[header_index + 1 :],
        start=header_index + 2,
    ):
        if all(value is None or not str(value).strip() for value in row):
            if holdings:
                break
            continue
        holding = _holding_from_row(row, columns, row_number=row_number)
        if holding.local_symbol in seen:
            raise HoldingsSchemaError(
                "Holdings workbook contains duplicate ticker "
                f"{holding.local_symbol!r} on rows "
                f"{seen[holding.local_symbol]} and {row_number}"
            )
        seen[holding.local_symbol] = row_number
        holdings.append(holding)

    if not holdings:
        raise ValueError("The workbook contained no recognisable holdings")

    file_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    effective_date = as_of_date_override or _as_of_date(rows[:header_index])
    if effective_date is None:
        raise ValueError(
            "Could not parse an 'As at' date; supply --as-of-date YYYY-MM-DD"
        )
    return HoldingsFile(
        path=path,
        as_of_date=effective_date,
        fund_name=fund_name,
        holdings=tuple(holdings),
        sha256=file_hash,
    )
